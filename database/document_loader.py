import os
import re
import hashlib
from PyPDF2 import PdfReader
from pptx import Presentation
import docx2txt
from config import Config
import traceback
import logging
from io import StringIO
from pdfminer.high_level import extract_text_to_fp
from pdfminer.layout import LAParams

logger = logging.getLogger(__name__)

def process_file(file_path):
    rel_path = os.path.relpath(file_path, Config.DATA_DIR)
    try:
        file_ext = os.path.splitext(file_path)[1].lower()
        if file_ext == '.pdf':
            text, metadata = process_pdf(file_path)
            return text, rel_path, metadata
        elif file_ext == '.pptx':
            text, metadata = process_ppt(file_path)
            return text, rel_path, metadata
        elif file_ext == '.docx':
            text, metadata = process_docx(file_path)
            return text, rel_path, metadata
        else:
            logger.warning(f"Unsupported format: {file_path}")
            return "", rel_path, {}
    except Exception as e:
        logger.error(f"Processing error: {str(e)}")
        return "", rel_path, {}

def process_pdf(file_path):
    text = ""
    metadata = {"pages": []}
    try:
        with open(file_path, 'rb') as f:
            pdf = PdfReader(f)
            for page_num in range(len(pdf.pages)):
                try:
                    page = pdf.pages[page_num]
                    page_text = page.extract_text() or extract_page_with_pdfminer(file_path, page_num)
                    if page_text:
                        text += f"<PAGE_START page={page_num+1}>\n{page_text}\n<PAGE_END>\n"
                        metadata["pages"].append({
                            "number": page_num+1,
                            "content": page_text[:500] + "..."
                        })
                except:
                    continue
        return text, metadata
    except Exception as e:
        logger.error(f"PDF error: {str(e)}")
        return "", metadata

def extract_page_with_pdfminer(file_path, page_number):
    output_string = StringIO()
    laparams = LAParams()
    try:
        with open(file_path, 'rb') as f:
            extract_text_to_fp(f, output_string, laparams=laparams, 
                              output_type='text', codec='utf-8', page_numbers=[page_number])
        return output_string.getvalue()
    except:
        return ""

def process_ppt(file_path):
    text = ""
    metadata = {"slides": []}
    try:
        prs = Presentation(file_path)
        for slide_num, slide in enumerate(prs.slides):
            slide_text = ""
            slide_title = slide.shapes.title.text.strip() if slide.shapes.title else ""
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip() and shape.text.strip() != slide_title:
                    slide_text += shape.text + " "
            if slide_text or slide_title:
                text += f"<SLIDE_START slide={slide_num+1} title='{slide_title}'>\n{slide_text}\n<SLIDE_END>\n"
                metadata["slides"].append({
                    "number": slide_num+1,
                    "title": slide_title,
                    "content": slide_text[:500] + "..." if slide_text else ""
                })
        return text, metadata
    except Exception as e:
        logger.error(f"PPT error: {str(e)}")
        return "", metadata

def process_docx(file_path):
    try:
        text = docx2txt.process(file_path)
        sections = []
        for section_num, paragraph in enumerate(text.split('\n\n')):
            if paragraph.strip():
                sections.append({
                    "number": section_num+1,
                    "content": paragraph[:500] + "..."
                })
        return text, {"sections": sections}
    except Exception as e:
        logger.error(f"DOCX error: {str(e)}")
        return "", {}

def chunk_text(text, source, chunk_size=1000, chunk_overlap=200):
    chunks = []
    current_chunk = ""
    current_page = 1
    current_slide = 1
    current_section = 1
    
    lines = text.split('\n')
    for line in lines:
        if line.startswith('<PAGE_START page='):
            current_page = int(re.search(r'page=(\d+)', line).group(1))
            continue
        elif line.strip() == '<PAGE_END>':
            continue
        if line.startswith('<SLIDE_START slide='):
            current_slide = int(re.search(r'slide=(\d+)', line).group(1))
            continue
        elif line.strip() == '<SLIDE_END>':
            continue
            
        current_chunk += line + '\n'
        if len(current_chunk) >= chunk_size:
            chunks.append({
                "text": current_chunk[:chunk_size],
                "metadata": {
                    "source": source,
                    "page": current_page,
                    "slide": current_slide,
                    "section": current_section
                }
            })
            current_chunk = current_chunk[chunk_size - chunk_overlap:]
            current_section += 1
    
    if current_chunk.strip():
        chunks.append({
            "text": current_chunk.strip(),
            "metadata": {
                "source": source,
                "page": current_page,
                "slide": current_slide,
                "section": current_section
            }
        })
    
    return chunks

def get_module_from_path(file_path):
    try:
        normalized_path = file_path.replace("\\", "/").lower()
        path_parts = normalized_path.split("/")
        for part in path_parts:
            if "module" in part:
                match = re.search(r"module(\d+)", part)
                if match:
                    return int(match.group(1))
        filename = os.path.basename(normalized_path)
        if "mod1" in filename: return 1
        if "mod2" in filename: return 2
        if "mod3" in filename: return 3
        if "mod4" in filename: return 4
        return 1
    except:
        return 1